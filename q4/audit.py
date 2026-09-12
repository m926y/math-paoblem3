"""Independent fourth-question physical, settlement, publication, and workbook audit."""
import numpy as np
from openpyxl import load_workbook
from microgrid.io import interval_ranges
from microgrid.timegrid import interval


def audit(records,summary,question,plans=None,workbook=None):
    checks={}
    arrays={k:np.array([r[k] for r in records]) for k in ('grid','grid0','charge','discharge','emergency','spill','energy','load','pv','price')}
    g,g0,c,r,e,s,E,L,V,p=[arrays[k] for k in arrays]
    checks['finite']=bool(all(np.isfinite(a).all() for a in arrays.values()))
    checks['shapes']=g.shape==(len(records),144) and E.shape==(len(records),145)
    checks['nonnegative']=bool(min(a.min() for a in (g,g0,c,r,e,s,L,V))>=-1e-7 and p.min()>0)
    residual=float(np.max(np.abs(g+V+r+e-L-c-s)))
    storage=float(np.max(np.abs(np.diff(E,axis=1)-.9*c+r/.9)))
    checks['balance']=residual<=1e-6;checks['storage']=storage<=1e-6
    # A deliberately fixed February state is allowed only at this one boundary.
    jump=np.abs(E[1:,0]-E[:-1,-1])
    if summary['initial_energy_policy']=='fixed_B_february' and len(jump)>30: jump[30]=0
    checks['continuity']=bool(not len(jump) or jump.max()<=1e-6)
    checks['bounds']=bool(E.min()>=1200-1e-7 and E.max()<=10800+1e-7 and max(c.max(),r.max())<=5000/6+1e-7)
    checks['mutual_exclusion']=bool(np.minimum(c,r).max()<=1e-6)
    checks['plan_residual']=summary['max_plan_residual_kwh']<=1e-6
    if question==2: checks['fixed_day_orders']=bool(np.max(np.abs(g-g0))<=1e-6)
    sel=slice(31,None) if len(records)>31 else slice(None)
    normal=np.sum(p[sel]*g[sel]);adjust=.5*np.sum(p[sel]*np.abs(g[sel]-g0[sel])) if question==3 else 0.
    urgent=5*np.sum(p[sel]*e[sel])
    for name,v in [('normal_cost_cny',normal),('adjustment_cost_cny',adjust),('emergency_cost_cny',urgent),('total_cost_cny',normal+adjust+urgent)]:
        checks[name]=bool(abs(float(v)-summary[name])<1e-5)
    if plans:
        index={(x['day'],x['issue']):x for x in plans}
        allowed=(0,) if question==2 else (0,6,12,18)
        checks['issue_times']=len(index)==len(records)*len(allowed) and all(k[1] in allowed for k in index)
        correct=True
        for day in range(len(records)):
            for issue in allowed:
                start=issue*6;take=144 if question==2 else 36;plan=index[day,issue]
                correct &= np.max(np.abs(g[day,start:start+take]-plan['grid'][:take]))<=1e-6
                correct &= abs(plan['energy_kwh']-E[day,start])<=1e-6
        checks['published_plan_execution']=bool(correct)
    if workbook:
        wb=load_workbook(workbook,read_only=True,data_only=True)
        selected=records[31:]
        checks['workbook_days']=len(selected)==334
        expected=['计划购电量']+(['调整购电量'] if question==3 else [])+['充放电量','紧急购电量']
        checks['sheets']=wb.sheetnames==expected
        for name in expected[:1+(question==3)]:
            rows=list(wb[name].values);actual=np.array([x[1:145] for x in rows[1:]],float)
            target=g0[31:] if name=='计划购电量' else g[31:]
            cost=(p[31:]*target).sum(1)
            if name=='调整购电量':cost+=.5*(p[31:]*np.abs(g[31:]-g0[31:])).sum(1)
            checks[name+'_dates']=[x[0].date() for x in rows[1:]]==[x['date'] for x in selected]
            checks[name+'_labels']=list(rows[0][1:145])==[interval(t) for t in range(144)]
            checks[name+'_values']=actual.shape==target.shape and bool(np.max(np.abs(actual-target))<=1e-6)
            checks[name+'_energy']=bool(np.max(np.abs(np.array([x[145] for x in rows[1:]])-target.sum(1)))<=1e-6)
            checks[name+'_fees']=bool(np.max(np.abs(np.array([x[146] for x in rows[1:]])-cost))<=1e-6)
        rows=list(wb['充放电量'].iter_rows(min_row=2,values_only=True))
        checks['storage_rows']=len(rows)==334*6
        checks['storage_values']=all(abs(row[2]-c[31+i//6,(i%6)*24:(i%6+1)*24].sum())<=1e-6 and
                                    abs(row[3]-r[31+i//6,(i%6)*24:(i%6+1)*24].sum())<=1e-6 for i,row in enumerate(rows))
        checks['storage_ends']=all(abs(rows[i*6][5]-E[i+31,0])<=1e-6 and abs(rows[i*6+1][5]-E[i+31,-1])<=1e-6 for i in range(334))
        observed=[];active=None
        for row in wb['紧急购电量'].iter_rows(min_row=2,values_only=True):
            if hasattr(row[0],'date'):active=row[0].date()
            if isinstance(row[2],(float,int)):observed.append((active,row[1],float(row[2])))
        expected_rows=[(rec['date'],interval(l,h+1),v) for rec in selected for l,h,v in interval_ranges(rec['emergency'])]
        checks['urgent_intervals']=len(observed)==len(expected_rows) and all(a[:2]==b[:2] and abs(a[2]-b[2])<=1e-6 for a,b in zip(observed,expected_rows))
        wb.close()
    result=dict(passed=all(checks.values()),checks=checks,max_balance_residual_kwh=residual,max_storage_residual_kwh=storage)
    if not result['passed']: raise AssertionError(result)
    return result
