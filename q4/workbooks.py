"""Fourth-question template writer adapted from the existing project writer.
Original microgrid/io.py remains unchanged; this copy accepts one price row per day.
"""
from pathlib import Path
from datetime import date,time
from copy import copy
import numpy as np
from openpyxl import load_workbook
from microgrid.timegrid import interval
from microgrid.io import interval_ranges
from microgrid.settlement import settle
T=144

def write_result2(template: Path, output: Path, dates: list[date], price: np.ndarray,
                  records: list[dict], question: int = 2) -> None:
    workbook = load_workbook(template)
    required = ["计划购电量", "充放电量", "紧急购电量"]
    if question == 3:
        required.insert(1,'调整购电量')
    if workbook.sheetnames != required:
        raise ValueError(f"result2模板工作表必须为：{required}")
    plan = workbook[required[0]]
    storage = workbook['充放电量']
    emergency = workbook['紧急购电量']
    if len(records) != len(dates):
        raise ValueError("result2记录数和日期数不一致")

    for row_no, (date_value, record) in enumerate(zip(dates, records), start=2):
        grid = np.asarray(record['grid0'] if question == 3 else record['grid'], dtype=float)
        plan.cell(row_no, 1, date_value).number_format = "yyyy-m-d"
        for t, value in enumerate(grid, start=2):
            plan.cell(row_no, t, float(value)).number_format = "0.0000"
        plan.cell(row_no, 146, float(np.sum(grid))).number_format = "0.0000"
        plan.cell(row_no, 147, float(np.dot(price[row_no-2], grid))).number_format = "0.0000"
        if question == 3:
            final = workbook['调整购电量']
            final.cell(row_no,1,date_value).number_format='yyyy-m-d'
            for t,value in enumerate(record['grid'],2):
                final.cell(row_no,t,float(value)).number_format='0.0000'
            final.cell(row_no,146,float(np.sum(record['grid']))).number_format='0.0000'
            fee=settle(price[row_no-2],record['grid'],record['emergency'],record['grid0'])
            final.cell(row_no,147,fee['normal_cost_cny']+fee['adjustment_cost_cny']).number_format='0.0000'

    for name in ('计划购电量','调整购电量') if question==3 else ('计划购电量',):
        for t in range(T): workbook[name].cell(1,t+2,interval(t))

    storage_styles = [
        [copy(storage.cell(row, col)._style) for col in range(1, 7)]
        for row in range(2, 8)
    ]
    if storage.max_row >= 2:
        storage.delete_rows(2, storage.max_row - 1)
    row_no = 2
    for date_value, record in zip(dates, records):
        charge = np.asarray(record["charge"], dtype=float)
        discharge = np.asarray(record["discharge"], dtype=float)
        soc = np.asarray(record['energy'], dtype=float)
        for block in range(6):
            row = storage_styles[block]
            storage.cell(row_no, 1, date_value if block == 0 else None)
            storage.cell(row_no, 2, f"{4*block}:00-{4*(block+1)}:00")
            storage.cell(row_no, 3, float(np.sum(charge[24*block:24*(block+1)])))
            storage.cell(row_no, 4, float(np.sum(discharge[24*block:24*(block+1)])))
            if block == 0:
                storage.cell(row_no, 5, time(0, 0))
                storage.cell(row_no, 6, float(soc[0]))
            elif block == 1:
                storage.cell(row_no, 5, "24:00")
                storage.cell(row_no, 6, float(soc[-1]))
            for col in range(1, 7):
                storage.cell(row_no, col)._style = copy(row[col-1])
            storage.cell(row_no, 1).number_format = "yyyy-m-d"
            for col in (3, 4, 6):
                storage.cell(row_no, col).number_format = "0.0000"
            row_no += 1

    emergency_styles = [copy(emergency.cell(2, col)._style) for col in range(1, 4)]
    if emergency.max_row >= 2:
        emergency.delete_rows(2, emergency.max_row - 1)
    row_no = 2
    for date_value, record in zip(dates, records):
        for left, right, amount in interval_ranges(record["emergency"]):
            emergency.cell(row_no, 1, date_value)
            emergency.cell(row_no, 2, interval(left,right+1))
            emergency.cell(row_no, 3, amount)
            for col in range(1, 4):
                emergency.cell(row_no, col)._style = copy(emergency_styles[col-1])
            emergency.cell(row_no, 1).number_format = "yyyy-m-d"
            emergency.cell(row_no, 3).number_format = "0.0000"
            row_no += 1
    if row_no == 2:
        emergency.cell(2, 1, "无紧急购电")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
