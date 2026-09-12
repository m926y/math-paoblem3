import argparse,json,sys,platform
from dataclasses import replace,asdict
from pathlib import Path
import numpy as np
import scipy
from openpyxl import load_workbook
from microgrid.io import attachment_paths,read_attachment1,read_attachment2,read_attachment3,file_sha256
from microgrid.forecast import ForecastEngine
from microgrid.reporting import save_json,csv_file
from microgrid.timegrid import validate_end_labels
from .model import PriceEngine,Settings,simulate
from .audit import audit
from .workbooks import write_result2
from .tables import paper_tables


def inputs(root):
    paths=attachment_paths(root)
    _,load0,pv0=read_attachment1(paths['data1'])
    dates,load,pv=read_attachment2(paths['data2'])
    f=read_attachment3(paths['data3'],dates)
    w=load_workbook(paths['data4'],read_only=True,data_only=True)
    try:
        rows=list(w.active.values)
        validate_end_labels(rows[0][1:])
        if [r[0].date() for r in rows[1:]]!=dates:raise ValueError('Price/date mismatch')
        p=np.array([r[1:] for r in rows[1:]],float)
        if p.shape!=(365,144):raise ValueError('Invalid annual price shape')
    finally:w.close()
    return ForecastEngine(load,pv,f,load0,pv0),PriceEngine(p),paths


def write_run(folder,records,plans,summary,question,paths,official=False):
    folder.mkdir(parents=True,exist_ok=True)
    keys=('grid','grid0','plan_charge','plan_discharge','charge','discharge','emergency','spill','energy','load','pv','price','price_prediction','load_prediction','pv_prediction','margin')
    np.savez_compressed(folder/'trace.npz',**{k:np.array([r[k] for r in records]) for k in keys})
    if plans:np.savez_compressed(folder/'plans.npz',**{k:np.array([r[k] for r in plans]) for k in plans[0]})
    csv_file(folder/'daily_summary.csv',['date','normal_cost_cny','adjustment_cost_cny','emergency_cost_cny','total_cost_cny','emergency_kwh','initial_kwh','terminal_kwh'],
             [[str(r['date']),r['normal_cost_cny'],r['adjustment_cost_cny'],r['emergency_cost_cny'],r['total_cost_cny'],float(r['emergency'].sum()),r['energy'][0],r['energy'][-1]] for r in records])
    wbpath=None
    if official:
        selected=records[31:]
        wbpath=folder/f'result4-{question}.xlsx'
        template=paths['attachment_dir']/'附件5'/f'result4-{question}.xlsx'
        write_result2(template,wbpath,[r['date'] for r in selected],np.array([r['price'] for r in selected]),selected,question)
        paper_tables(folder,selected,question,None)
        summary['template_sha256']=file_sha256(template)
    verification=audit(records,summary,question,plans,wbpath)
    summary['validation_passed']=verification['passed']
    summary['input_sha256']={f'data{i}':file_sha256(paths[f'data{i}']) for i in range(1,5)}
    summary['environment']={'python':sys.version,'numpy':np.__version__,'scipy':scipy.__version__,'platform':platform.platform()}
    summary['source_sha256']={f.name:file_sha256(f) for f in Path(__file__).parent.glob('*.py')}
    summary['method_note']='causal unknown future prices; shared storage; final deviation from midnight plan; initial day zero orders'
    save_json(folder/'summary.json',summary);save_json(folder/'validation.json',verification)
    save_json(folder/'frozen_numbers.json',summary)


def event(folder,payload):
    folder.mkdir(parents=True,exist_ok=True)
    with (folder/'events.jsonl').open('a',encoding='utf8') as f:f.write(json.dumps(payload,ensure_ascii=False,allow_nan=False)+'\n')


def select_b(engine,prices,q,folder):
    candidates=[]
    for method in ('persistence','weighted'):
        for alpha in (.5,.65,.8):
            cfg=Settings(price_method=method,quantile=alpha)
            try:
                rec,plans,s=simulate(engine,prices,q,cfg,days=31)
                audit(rec,s,q)
            except Exception as exc:
                event(folder,dict(stage='B_selection',question=q,config=asdict(cfg),failed=str(exc)));raise
            candidates.append(s);event(folder,dict(stage='B_selection',question=q,summary=s))
            print(f'Q4-{q} JAN B {method} {alpha}: {s["total_cost_cny"]:.4f}',flush=True)
    minimum=min(s['total_cost_cny'] for s in candidates)
    chosen=min((s for s in candidates if s['total_cost_cny']<=minimum+.01),key=lambda s:(s['config']['price_method']!='persistence',s['config']['quantile']))
    save_json(folder/f'q{q}_B_selection.json',dict(selection_period='2025-01-01/2025-01-31',evaluation_period='2025-02-01/2025-12-31',chosen=chosen['config'],candidates=candidates))
    return Settings(**chosen['config'])


def select_c(engine,prices,q,base,folder):
    rows=[]
    for name,risk in (('C_fixed',0),('C0',0),('C15',.15)):
        cfg=replace(base,variant=name,risk_weight=risk)
        rec,_,s=simulate(engine,prices,q,cfg,days=31);audit(rec,s,q)
        rows.append(s);event(folder,dict(stage='C_selection',question=q,summary=s))
        print(f'Q4-{q} JAN {name}: {s["total_cost_cny"]:.4f}',flush=True)
    candidates=[r for r in rows if r['config']['variant']!='C_fixed']
    low=min(r['total_cost_cny'] for r in candidates)
    winner=min((r for r in candidates if r['total_cost_cny']<=low+.01),key=lambda r:r['config']['risk_weight'])
    save_json(folder/f'q{q}_C_selection.json',dict(chosen=winner['config'],candidates=rows,selection_period='January only'))
    return winner['config']['variant']


def run(root,question=0,experiments=False):
    engine,prices,paths=inputs(root)
    destination=root/'output/q4_experiments'
    for q in ((2,3) if not question else (question,)):
        cfg=select_b(engine,prices,q,destination)
        chosen_c=select_c(engine,prices,q,cfg,destination) if experiments else None
        rec,plans,summary=simulate(engine,prices,q,cfg,snapshots=True,progress=True)
        write_run(root/f'output/q4-{q}',rec,plans,summary,q,paths,official=True)
        print(f'Q4-{q} B total {summary["total_cost_cny"]:.4f}; validated',flush=True)
        if experiments:
            for name,risk in (('C_fixed',0),('C0',0),('C15',.15)):
                candidate=replace(cfg,variant=name,risk_weight=risk)
                try:
                    rs,ps,s=simulate(engine,prices,q,candidate,snapshots=True,progress=True)
                    s['selected_on_january']=name==chosen_c
                    write_run(destination/f'q{q}_{name}',rs,ps,s,q,paths)
                    event(destination,dict(stage='full_year',question=q,summary=s))
                except Exception as exc:
                    event(destination,dict(stage='full_year',question=q,variant=name,failed=str(exc)));raise
                print(f'Q4-{q} {name} total {s["total_cost_cny"]:.4f}',flush=True)
            # Only the January-selected joint C gets an extra matched-initial-state comparison.
            candidate=replace(cfg,variant=chosen_c,risk_weight=.15 if chosen_c=='C15' else 0)
            rs,ps,s=simulate(engine,prices,q,candidate,february_energy=summary['february_initial_energy_kwh'],progress=True)
            write_run(destination/f'q{q}_C_matched_initial',rs,ps,s,q,paths)
            event(destination,dict(stage='matched_initial',question=q,summary=s))


def main(root):
    parser=argparse.ArgumentParser(description='独立第四问：方案B主结果与方案C评估，不改变前三问')
    parser.add_argument('--question',type=int,choices=(2,3),default=0)
    parser.add_argument('--experiments',action='store_true')
    args=parser.parse_args()
    run(Path(root),args.question,args.experiments)
