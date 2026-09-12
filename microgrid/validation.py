"""从逐段轨迹独立重算，校验不能只信任摘要。"""
import json
from datetime import date,timedelta
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from .reporting import save_json
from .io import file_sha256
from .timegrid import interval


def validate_output(folder,question):
    folder=Path(folder)
    data=np.load(folder/'trace.npz',allow_pickle=False)
    summary=json.loads((folder/'summary.json').read_text(encoding='utf-8'))
    checks=[]
    def check(name,ok):
        checks.append(dict(name=name,passed=bool(ok)))
    g,c,r,e,s,E,L,V=[data[k] for k in ('grid','charge','discharge','emergency','spill','energy','load','pv')]
    p=data['price'];g0=data['grid0']; sel=slice(31,None)
    check('365天144段与145储能点',g.shape==(365,144) and E.shape==(365,145))
    check('全部数据有限',all(np.isfinite(data[k]).all() for k in data.files))
    check('非负电量',min(x.min() for x in (g,c,r,e,s,L,V))>=-1e-7)
    check('供需平衡',np.max(np.abs(g+V+r+e-L-c-s))<=1e-6)
    check('储能递推',np.max(np.abs(np.diff(E,axis=1)-.9*c+r/.9))<=1e-6)
    check('跨日状态连续',np.max(np.abs(E[1:,0]-E[:-1,-1]))<=1e-6)
    check('物理边界',E.min()>=1200-1e-7 and E.max()<=10800+1e-7 and max(c.max(),r.max())<=5000/6+1e-7)
    check('充放电互斥',np.minimum(c,r).max()<=1e-6)
    normal=(g[sel]*p).sum();adjust=.5*(np.abs(g[sel]-g0[sel])*p).sum() if question==3 else 0.
    urgent=5*(e[sel]*p).sum()
    for name,value in (('normal_cost_cny',normal),('adjustment_cost_cny',adjust),('emergency_cost_cny',urgent),('total_cost_cny',normal+adjust+urgent)):
        check('独立核算 '+name,abs(value-summary[name])<1e-4)
    wb=load_workbook(folder/f'result{question}.xlsx',read_only=True,data_only=True)
    for name,array in [('计划购电量',g0 if question==3 else g)]+([('调整购电量',g)] if question==3 else []):
        sheet=wb[name]
        rows=list(sheet.iter_rows(values_only=True))
        check(name+'时段',list(rows[0][1:145])==[interval(t) for t in range(144)])
        check(name+'天数',len(rows)==335)
        check(name+'日期完整',
              [row[0].date() if hasattr(row[0],'date') else row[0] for row in rows[1:]]==
              [date(2025,2,1)+timedelta(days=i) for i in range(334)])
        values=np.array([row[1:145] for row in rows[1:]],float)
        check(name+'逐段对应轨迹',np.max(np.abs(values-array[sel]))<=1e-6)
        costs=np.array([row[146] for row in rows[1:]])
        want=(array[sel]*p).sum(axis=1)
        if name=='调整购电量':want+=.5*(np.abs(g[sel]-g0[sel])*p).sum(axis=1)
        check(name+'逐日费用',np.max(np.abs(costs-want))<=1e-6)
    storage=list(wb['充放电量'].iter_rows(min_row=2,values_only=True))
    check('充放电表完整',len(storage)==334*6)
    check('储能表按4小时块对应',all(abs(row[2]-c[31+i//6,(i%6)*24:(i%6+1)*24].sum())<1e-6 and
          abs(row[3]-r[31+i//6,(i%6)*24:(i%6+1)*24].sum())<1e-6 for i,row in enumerate(storage)))
    check('储能表日初日末电量',all(abs(storage[i*6][5]-E[31+i,0])<1e-6 and
          abs(storage[i*6+1][5]-E[31+i,-1])<1e-6 for i in range(334)))
    amount=sum(row[2] for row in wb['紧急购电量'].iter_rows(min_row=2,values_only=True) if isinstance(row[2],(float,int)))
    check('紧急购电表总量',abs(amount-e[sel].sum())<1e-5)
    wb.close()
    if (folder/'plans.npz').exists():
        plans=np.load(folder/'plans.npz',allow_pickle=False)
        check('计划仅规定发布时间',set(plans['issue']).issubset({0,6,12,18}))
        last={(int(day),int(issue)):i for i,(day,issue) in enumerate(zip(plans['day'],plans['issue']))}
        # 每个实际购电段来自该段以前最后一份已发布计划。
        valid=True
        for day in range(365):
            issued=sorted(k[1] for k in last if k[0]==day)
            for t in range(144):
                hour=max(h for h in issued if h*6<=t)
                if abs(g[day,t]-plans['grid'][last[day,hour],t-hour*6])>1e-6:valid=False
        check('执行购电对应最近已发布计划',valid)
    result={'passed':all(c['passed'] for c in checks),'checks':checks,
            'validator_sha256':file_sha256(Path(__file__)),
            'validated_files_sha256':{name:file_sha256(folder/name) for name in
                ('trace.npz',f'result{question}.xlsx','summary.json')}}
    save_json(folder/'validation.json',result)
    if not result['passed']:
        raise AssertionError([c['name'] for c in checks if not c['passed']])
    return result
