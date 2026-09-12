"""项目统一输入、路径和官方结果模板读写。"""

from __future__ import annotations

import hashlib
import math
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


T = 144
EXPECTED_DATES = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
ISSUES = ("0:00", "6:00", "12:00", "18:00")


def attachment_paths(base_dir: Path, attachment_dir: Path | None = None) -> dict[str, Path]:
    directory = (attachment_dir or (base_dir / "附件")).resolve()
    paths = {
        "attachment_dir": directory,
        "data1": directory / "附件1.xlsx",
        "data2": directory / "附件2.xlsx",
        "data3": directory / "附件3.xlsx",
        "data4": directory / "附件4.xlsx",
        "template1": directory / "附件5" / "result1.xlsx",
        "template2": directory / "附件5" / "result2.xlsx",
        "template3": directory / "附件5" / "result3.xlsx",
    }
    required = ["data1", "data2", "data3", "template1", "template2", "template3"]
    missing = [f"{key}: {paths[key]}" for key in required if not paths[key].is_file()]
    if missing:
        raise FileNotFoundError("缺少输入文件：\n" + "\n".join(missing))
    return paths


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _date_value(value: object, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field}不是有效日期：{value!r}") from exc


def _number(value: object, field: str, *, nonnegative: bool = True) -> float:
    if value is None or value == "" or isinstance(value, bool):
        raise ValueError(f"{field}缺失或不是数值")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field}不是数值：{value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{field}包含NaN或无穷值")
    if nonnegative and result < 0:
        raise ValueError(f"{field}不能为负：{result}")
    return result


def read_attachment1(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(min_row=2, max_row=145, max_col=4, values_only=True))
        if len(rows) != T or sheet.max_row != T + 1 or sheet.max_column < 4:
            raise ValueError("附件1必须为1行表头和144行数据，且至少包含4列")
        price = np.array([_number(row[1], f"附件1第{i+2}行电价") for i, row in enumerate(rows)])
        load = np.array([_number(row[2], f"附件1第{i+2}行负荷") for i, row in enumerate(rows)])
        pv = np.array([_number(row[3], f"附件1第{i+2}行光伏") for i, row in enumerate(rows)])
        return price, load, pv
    finally:
        workbook.close()


def _read_daily_sheet(sheet, label: str) -> tuple[list[date], np.ndarray]:
    rows = list(sheet.iter_rows(min_row=2, max_col=T + 1, values_only=True))
    if len(rows) != 365 or sheet.max_column < T + 1:
        raise ValueError(f"{label}必须包含365天且每天144个十分钟点")
    dates: list[date] = []
    matrix = np.empty((365, T), dtype=float)
    for i, row in enumerate(rows):
        dates.append(_date_value(row[0], f"{label}第{i+2}行日期"))
        if len(row) < T + 1:
            raise ValueError(f"{label}第{i+2}行数据不足144个")
        for t, value in enumerate(row[1:T+1]):
            matrix[i, t] = _number(value, f"{label}第{i+2}行第{t+1}时段")
    if dates != EXPECTED_DATES:
        raise ValueError(f"{label}日期必须无重复、无缺失并连续覆盖2025-01-01至2025-12-31")
    return dates, matrix


def read_attachment2(path: Path) -> tuple[list[date], np.ndarray, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        required = ("小区负载", "光伏发电实际功率")
        missing = [name for name in required if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"附件2缺少工作表：{missing}")
        load_dates, load = _read_daily_sheet(workbook[required[0]], required[0])
        pv_dates, pv = _read_daily_sheet(workbook[required[1]], required[1])
        if load_dates != pv_dates:
            raise ValueError("附件2负荷与光伏工作表日期不一致")
        return load_dates, load, pv
    finally:
        workbook.close()


def _issue_label(value: object) -> str:
    if isinstance(value, time):
        return f"{value.hour}:{value.minute:02d}"
    if isinstance(value, datetime):
        return f"{value.hour}:{value.minute:02d}"
    text = str(value).strip()
    if text.startswith("0") and text not in ("0", "0:00"):
        text = text.lstrip("0") or "0"
        if text.startswith(":"):
            text = "0" + text
    return text


def read_attachment3(path: Path, dates: list[date]) -> np.ndarray:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        date_index = {value: i for i, value in enumerate(dates)}
        issue_index = {value: i for i, value in enumerate(ISSUES)}
        forecasts = np.full((len(dates), len(ISSUES), 24), np.nan, dtype=float)
        seen: set[tuple[date, str]] = set()
        current_date: date | None = None
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, max_col=26, values_only=True), start=2):
            if all(value in (None, "") for value in row):
                continue
            if row[0] not in (None, ""):
                current_date = _date_value(row[0], f"附件3第{row_no}行日期")
            if current_date is None:
                raise ValueError(f"附件3第{row_no}行缺少可继承的日期")
            if current_date not in date_index:
                raise ValueError(f"附件3出现未知日期：{current_date}")
            issue = _issue_label(row[1])
            if issue not in issue_index:
                raise ValueError(f"附件3第{row_no}行发布时间无效：{row[1]!r}")
            key = (current_date, issue)
            if key in seen:
                raise ValueError(f"附件3存在重复记录：{current_date} {issue}")
            seen.add(key)
            if len(row) < 26:
                raise ValueError(f"附件3第{row_no}行预测值不足24个")
            forecasts[date_index[current_date], issue_index[issue]] = [
                _number(value, f"附件3第{row_no}行提前{h+1}小时")
                for h, value in enumerate(row[2:26])
            ]
        expected = {(d, issue) for d in dates for issue in ISSUES}
        missing = sorted(expected - seen)
        if missing:
            sample = ", ".join(f"{d} {issue}" for d, issue in missing[:8])
            raise ValueError(f"附件3缺少{len(missing)}条日期×发布时间记录，例如：{sample}")
        if not np.isfinite(forecasts).all():
            raise ValueError("附件3仍包含未读取或无效的预测值")
        return forecasts
    finally:
        workbook.close()


def _copy_style(source, target) -> None:
    target._style = copy(source._style)
    if source.has_style:
        target.number_format = source.number_format


def _clear_values(sheet, min_row: int = 2) -> None:
    for row in sheet.iter_rows(min_row=min_row):
        for cell in row:
            cell.value = None


def write_result1(template: Path, output: Path, grid: np.ndarray, charge: np.ndarray,
                  discharge: np.ndarray, soc: np.ndarray) -> None:
    workbook = load_workbook(template)
    if workbook.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError("result1模板工作表结构不符合题目")
    plan = workbook["计划购电量"]
    storage = workbook["充放电量"]
    for t, value in enumerate(np.asarray(grid), start=2):
        plan.cell(t, 2, float(value)).number_format = "0.0000"
    for block in range(6):
        left, right = block * 24, (block + 1) * 24
        storage.cell(block + 2, 2, float(np.sum(charge[left:right]))).number_format = "0.0000"
        storage.cell(block + 2, 3, float(np.sum(discharge[left:right]))).number_format = "0.0000"
    storage.cell(2, 5, float(soc[0])).number_format = "0.0000"
    storage.cell(3, 5, float(soc[-1])).number_format = "0.0000"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def interval_ranges(values: np.ndarray, tolerance: float = 1e-8):
    values = np.asarray(values, dtype=float)
    i = 0
    while i < len(values):
        if values[i] <= tolerance:
            i += 1
            continue
        j = i
        while j + 1 < len(values) and values[j + 1] > tolerance:
            j += 1
        yield i, j, float(np.sum(values[i:j+1]))
        i = j + 1


def clock_label(index: int) -> str:
    minutes = index * 10
    suffix = "+1" if minutes >= 1440 else ""
    minutes %= 1440
    return f"{minutes // 60}:{minutes % 60:02d}{suffix}"


def write_result2(template: Path, output: Path, dates: list[date], price: np.ndarray,
                  records: list[dict]) -> None:
    workbook = load_workbook(template)
    required = ["计划购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != required:
        raise ValueError(f"result2模板工作表必须为：{required}")
    plan = workbook[required[0]]
    storage = workbook[required[1]]
    emergency = workbook[required[2]]
    if len(records) != len(dates):
        raise ValueError("result2记录数和日期数不一致")

    for row_no, (date_value, record) in enumerate(zip(dates, records), start=2):
        grid = np.asarray(record["grid"], dtype=float)
        plan.cell(row_no, 1, date_value).number_format = "yyyy-m-d"
        for t, value in enumerate(grid, start=2):
            plan.cell(row_no, t, float(value)).number_format = "0.0000"
        plan.cell(row_no, 146, float(np.sum(grid))).number_format = "0.0000"
        plan.cell(row_no, 147, float(np.dot(price, grid))).number_format = "0.0000"

    storage_styles = [
        [copy(storage.cell(row, col)._style) for col in range(1, 7)]
        for row in range(2, 8)
    ]
    if storage.max_row >= 2:
        storage.delete_rows(2, storage.max_row - 1)
    row_no = 2
    for date_value, record in zip(dates, records):
        charge = np.asarray(record["charge"], dtype=float)
        discharge = np.asarray(record["discharge"], dtype=float)
        soc = np.asarray(record["soc"], dtype=float)
        for block in range(6):
            row = storage_styles[block]
            storage.cell(row_no, 1, date_value if block == 0 else None)
            storage.cell(row_no, 2, f"{4*block}:00-{4*(block+1)}:00")
            storage.cell(row_no, 3, float(np.sum(charge[24*block:24*(block+1)])))
            storage.cell(row_no, 4, float(np.sum(discharge[24*block:24*(block+1)])))
            if block == 0:
                storage.cell(row_no, 5, time(0, 0))
                storage.cell(row_no, 6, float(soc[0]))
            elif block == 1:
                storage.cell(row_no, 5, "24:00")
                storage.cell(row_no, 6, float(soc[-1]))
            for col in range(1, 7):
                storage.cell(row_no, col)._style = copy(row[col-1])
            storage.cell(row_no, 1).number_format = "yyyy-m-d"
            for col in (3, 4, 6):
                storage.cell(row_no, col).number_format = "0.0000"
            row_no += 1

    emergency_styles = [copy(emergency.cell(2, col)._style) for col in range(1, 4)]
    if emergency.max_row >= 2:
        emergency.delete_rows(2, emergency.max_row - 1)
    row_no = 2
    for date_value, record in zip(dates, records):
        for left, right, amount in interval_ranges(record["emergency"]):
            emergency.cell(row_no, 1, date_value)
            emergency.cell(row_no, 2, f"{clock_label(left)}-{clock_label(right+1)}")
            emergency.cell(row_no, 3, amount)
            for col in range(1, 4):
                emergency.cell(row_no, col)._style = copy(emergency_styles[col-1])
            emergency.cell(row_no, 1).number_format = "yyyy-m-d"
            emergency.cell(row_no, 3).number_format = "0.0000"
            row_no += 1
    if row_no == 2:
        emergency.cell(2, 1, "无紧急购电")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
