"""统一验证第一至三问的输入、结果表、约束摘要和模型版本对照。"""

from __future__ import annotations

import argparse
import csv
import json
import math
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from datetime import date, timedelta
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

from project_io import (
    attachment_paths,
    file_sha256,
    read_attachment1,
    read_attachment2,
    read_attachment3,
)
from q3_dispatch import DispatchParams, execute_causal_block


ROOT = Path(__file__).resolve().parent
EXPECTED_COSTS = {
    "q1": 35126.9485892896,
    "q2": 18698663.823237758,
    "q3": 17349812.52338954,
}
SCENARIOS = {
    "q1_engineering_baseline": ROOT / "output" / "engineering_baseline" / "q1",
    "q2_engineering_baseline": ROOT / "output" / "engineering_baseline" / "q2",
    "q2_causal_zero": ROOT / "output" / "model_versions" / "q2_causal_zero",
    "q2_causal_baseline": ROOT / "output" / "model_versions" / "q2_causal_baseline",
    "q2_weekday_optimized": ROOT / "output" / "q2",
    "q3_update_0": ROOT / "output" / "model_versions" / "q3_update_0",
    "q3_update_0_6": ROOT / "output" / "model_versions" / "q3_update_0_6",
    "q3_update_0_6_12": ROOT / "output" / "model_versions" / "q3_update_0_6_12",
    "q3_update_full_no_terminal": ROOT / "output" / "model_versions" / "q3_update_full_no_terminal",
    "q3_update_full_terminal": ROOT / "output" / "model_versions" / "q3_update_full_terminal",
}


class Audit:
    def __init__(self) -> None:
        self.checks: list[dict] = []

    def check(self, name: str, condition: bool, evidence: object) -> None:
        self.checks.append({"name": name, "passed": bool(condition), "evidence": evidence})

    def close(self, name: str, actual: float, expected: float, tolerance: float) -> None:
        self.check(name, abs(actual - expected) <= tolerance, {
            "actual": actual, "expected": expected, "absolute_difference": abs(actual - expected),
            "tolerance": tolerance,
        })

    @property
    def passed(self) -> bool:
        return all(item["passed"] for item in self.checks)


def _run_all() -> None:
    commands = [
        [sys.executable, "run_q1.py", "--output-dir", "output/engineering_baseline/q1"],
        [sys.executable, "run_q2_legacy.py", "--output-dir", "output/engineering_baseline/q2"],
        [sys.executable, "q3_model.py", "--output", "output/engineering_baseline/q3/result3.xlsx",
         "--summary", "output/engineering_baseline/q3/summary.json"],
        [sys.executable, "run_q2_causal.py", "--cold-start", "zero",
         "--output-dir", "output/model_versions/q2_causal_zero"],
        [sys.executable, "run_q2_causal.py", "--cold-start", "baseline",
         "--output-dir", "output/model_versions/q2_causal_baseline"],
        [sys.executable, "run_q2.py"],
    ]
    for label, hours in (("q3_update_0", "0"), ("q3_update_0_6", "0,6"),
                         ("q3_update_0_6_12", "0,6,12"),
                         ("q3_update_full_no_terminal", "0,6,12,18")):
        commands.append([
            sys.executable, "q3_model.py", "--update-hours", hours,
            "--output", f"output/model_versions/{label}/result3.xlsx",
            "--summary", f"output/model_versions/{label}/summary.json",
        ])
    commands.append([
        sys.executable, "q3_model.py", "--update-hours", "0,6,12,18",
        "--terminal-value-factor", "0.9",
        "--output", "output/model_versions/q3_update_full_terminal/result3.xlsx",
        "--summary", "output/model_versions/q3_update_full_terminal/summary.json",
    ])
    for command in commands:
        subprocess.run(command, cwd=ROOT, check=True)


def _expect_rejection(audit: Audit, name: str, action) -> None:
    try:
        action()
    except (ValueError, FileNotFoundError) as exc:
        audit.check(name, True, str(exc))
    else:
        audit.check(name, False, "无异常，错误输入被接受")


def _input_failure_tests(audit: Audit, data2: Path, data3: Path, dates: list[date]) -> None:
    with tempfile.TemporaryDirectory(prefix="model_validation_") as temp_name:
        temp_dir = Path(temp_name)

        missing = temp_dir / "forecast_missing.xlsx"
        shutil.copy2(data3, missing)
        workbook = load_workbook(missing)
        workbook.active.delete_rows(2, 1)
        workbook.save(missing)
        _expect_rejection(audit, "附件3缺失日期×发布时间记录会被拒绝",
                          lambda: read_attachment3(missing, dates))

        duplicate = temp_dir / "forecast_duplicate.xlsx"
        shutil.copy2(data3, duplicate)
        workbook = load_workbook(duplicate)
        sheet = workbook.active
        sheet.cell(3, 2, sheet.cell(2, 2).value)
        workbook.save(duplicate)
        _expect_rejection(audit, "附件3重复日期×发布时间记录会被拒绝",
                          lambda: read_attachment3(duplicate, dates))

        blank = temp_dir / "forecast_blank.xlsx"
        shutil.copy2(data3, blank)
        workbook = load_workbook(blank)
        workbook.active.cell(2, 3).value = None
        workbook.save(blank)
        _expect_rejection(audit, "附件3空预测值会被拒绝",
                          lambda: read_attachment3(blank, dates))

        bad_daily = temp_dir / "daily_blank.xlsx"
        shutil.copy2(data2, bad_daily)
        workbook = load_workbook(bad_daily)
        workbook["小区负载"].cell(2, 2).value = None
        workbook.save(bad_daily)
        _expect_rejection(audit, "附件2空负荷值会被拒绝",
                          lambda: read_attachment2(bad_daily))


def _numeric_values(sheet, row: int, left: int, right: int) -> np.ndarray:
    values = np.array([sheet.cell(row, col).value for col in range(left, right + 1)], dtype=float)
    if not np.isfinite(values).all() or np.any(values < -1e-8):
        raise ValueError(f"{sheet.title}第{row}行包含非法数值")
    return values


def _validate_q1_book(audit: Audit, folder: Path, summary: dict) -> None:
    path = folder / "result1.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    audit.check("Q1官方模板工作表", workbook.sheetnames == ["计划购电量", "充放电量"], workbook.sheetnames)
    sheet = workbook["计划购电量"]
    values = np.array([row[0] for row in sheet.iter_rows(
        min_row=2, max_row=145, min_col=2, max_col=2, values_only=True)], dtype=float)
    audit.close("Q1结果表购电量汇总", float(values.sum()), summary["grid_energy_kwh"], 0.01)
    workbook.close()


def _validate_plan_sheet(audit: Audit, sheet, price: np.ndarray, prefix: str,
                         expected_energy: float, expected_cost: float,
                         adjusted: bool = False, reference=None) -> None:
    dates = [date(2025, 2, 1) + timedelta(days=i) for i in range(334)]
    actual_dates = []
    total_energy = 0.0
    total_fee = 0.0
    rows = sheet.iter_rows(min_row=2, max_row=335, min_col=1, max_col=147, values_only=True)
    for row_no, (expected_date, row) in enumerate(zip(dates, rows), start=2):
        value = row[0]
        actual_date = value.date() if hasattr(value, "date") else value
        actual_dates.append(actual_date)
        series = np.asarray(row[1:145], dtype=float)
        if not np.isfinite(series).all() or np.any(series < -1e-8):
            audit.check(f"{prefix}数值有效", False, {"row": row_no})
            return
        total_energy += float(series.sum())
        total_fee += float(row[146])
        if abs(float(row[145]) - float(series.sum())) > 0.02:
            audit.check(f"{prefix}逐日总量", False, {"row": row_no})
            break
    else:
        audit.check(f"{prefix}逐日总量", True, "334行逐日总量一致")
    audit.check(f"{prefix}日期覆盖", actual_dates == dates, {
        "first": str(actual_dates[0]), "last": str(actual_dates[-1]), "rows": len(actual_dates)
    })
    audit.close(f"{prefix}电量汇总", total_energy, expected_energy, 0.05)
    audit.close(f"{prefix}费用汇总", total_fee, expected_cost, 0.05)


def _validate_q2_book(audit: Audit, folder: Path, summary: dict, price: np.ndarray, label: str) -> None:
    path = folder / "result2.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    expected = ["计划购电量", "充放电量", "紧急购电量"]
    audit.check(f"{label}官方模板工作表", workbook.sheetnames == expected, workbook.sheetnames)
    _validate_plan_sheet(audit, workbook["计划购电量"], price, f"{label}计划购电",
                         summary["normal_energy_kwh"], summary["normal_cost_cny"])
    audit.check(f"{label}充放电表行数", workbook["充放电量"].max_row == 2005,
                workbook["充放电量"].max_row)
    emergency_sum = sum(float(row[0] or 0.0) for row in workbook["紧急购电量"].iter_rows(
        min_row=2, min_col=3, max_col=3, values_only=True))
    audit.close(f"{label}紧急购电量汇总", emergency_sum, summary["emergency_energy_kwh"], 0.05)
    workbook.close()


def _validate_q3_book(audit: Audit, folder: Path, summary: dict, price: np.ndarray, label: str) -> None:
    path = folder / "result3.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    expected = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    audit.check(f"{label}官方模板工作表", workbook.sheetnames == expected, workbook.sheetnames)
    _validate_plan_sheet(audit, workbook["计划购电量"], price, f"{label}0点计划",
                         summary["plan_energy_kwh"], summary["plan_cost_cny"])
    _validate_plan_sheet(audit, workbook["调整购电量"], price, f"{label}最终正常购电",
                         summary["final_normal_energy_kwh"],
                         summary["final_normal_cost_cny"] + summary["adjustment_cost_cny"])
    audit.check(f"{label}充放电表行数", workbook["充放电量"].max_row == 2005,
                workbook["充放电量"].max_row)
    emergency_sum = sum(float(row[0] or 0.0) for row in workbook["紧急购电量"].iter_rows(
        min_row=2, min_col=3, max_col=3, values_only=True))
    audit.close(f"{label}紧急购电量汇总", emergency_sum, summary["emergency_energy_kwh"], 0.05)
    workbook.close()


def _constraint_checks(audit: Audit, label: str, summary: dict) -> None:
    for key in (
        "max_balance_residual_kwh", "max_plan_balance_residual_kwh",
        "max_plan_soc_residual_kwh", "max_actual_balance_residual_kwh",
        "max_actual_soc_residual_kwh", "max_charge_discharge_overlap_kwh",
        "max_simultaneous_charge_discharge_kwh",
    ):
        if key in summary:
            audit.check(f"{label} {key}", abs(float(summary[key])) <= 1e-6, summary[key])
    if "soc_min_kwh" in summary:
        audit.check(f"{label} SOC下界", float(summary["soc_min_kwh"]) >= 1200.0 - 1e-5,
                    summary["soc_min_kwh"])
        audit.check(f"{label} SOC上界", float(summary["soc_max_kwh"]) <= 10800.0 + 1e-5,
                    summary["soc_max_kwh"])


def _stress_test(audit: Audit) -> None:
    params = DispatchParams()
    zeros = np.zeros(12)
    shortage = execute_causal_block(zeros, zeros, zeros, np.full(12, 100000.0), zeros,
                                    params.e_initial, params)
    surplus = execute_causal_block(zeros, zeros, zeros, zeros, np.full(12, 100000.0),
                                   params.e_initial, params)
    evidence = {
        "shortage_balance_max": float(shortage["balance_residual_max"]),
        "surplus_balance_max": float(surplus["balance_residual_max"]),
        "shortage_emergency_kwh": float(np.sum(shortage["emergency"])),
        "surplus_spill_kwh": float(np.sum(surplus["spill"])),
    }
    audit.check("极端缺电与光伏过剩压力测试", evidence["shortage_balance_max"] <= 1e-8
                and evidence["surplus_balance_max"] <= 1e-8
                and evidence["shortage_emergency_kwh"] > 0 and evidence["surplus_spill_kwh"] > 0,
                evidence)


def _load_summaries(audit: Audit) -> dict[str, dict]:
    summaries = {}
    for label, folder in SCENARIOS.items():
        path = folder / "summary.json"
        if not path.is_file():
            audit.check(f"{label}摘要存在", False, str(path))
            continue
        summaries[label] = json.loads(path.read_text(encoding="utf-8"))
        audit.check(f"{label}摘要存在", True, str(path.relative_to(ROOT)))
    return summaries


def _write_comparison(summaries: dict[str, dict]) -> list[dict]:
    rows = []
    for label, summary in summaries.items():
        rows.append({
            "scenario": label,
            "model_version": summary.get("model_version", summary.get("model", "")),
            "update_hours": ",".join(map(str, summary.get("update_hours", []))),
            "terminal_value_factor": summary.get("terminal_value_factor", ""),
            "february_initial_soc_kwh": summary.get("february_initial_soc_kwh", ""),
            "normal_or_final_cost_cny": summary.get("normal_cost_cny", summary.get("final_normal_cost_cny", "")),
            "adjustment_cost_cny": summary.get("adjustment_cost_cny", ""),
            "emergency_cost_cny": summary.get("emergency_cost_cny", ""),
            "emergency_energy_kwh": summary.get("emergency_energy_kwh", ""),
            "total_cost_cny": summary.get("total_cost_cny", summary.get("objective_cost_cny", "")),
            "january_warmup_cost_cny": summary.get("january_warmup_cost_cny", ""),
        })
    destination = ROOT / "output" / "model_versions"
    destination.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0])
    with (destination / "model_comparison.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    (destination / "model_comparison.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "模型版本对照"
    sheet.append(fields)
    for row in rows:
        sheet.append([row[field] for field in fields])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    widths = [34, 40, 18, 22, 24, 26, 24, 24, 24, 24, 26]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=len(fields)):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000"
    workbook.save(destination / "model_comparison.xlsx")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment-dir", type=Path, default=ROOT / "附件")
    parser.add_argument("--rerun", action="store_true", help="先重新执行全部第一至三问及对照实验")
    args = parser.parse_args()
    if args.rerun:
        _run_all()

    audit = Audit()
    paths = attachment_paths(ROOT, args.attachment_dir)
    price, load1, pv1 = read_attachment1(paths["data1"])
    dates, load2, pv2 = read_attachment2(paths["data2"])
    forecasts = read_attachment3(paths["data3"], dates)
    audit.check("输入形状", price.shape == load1.shape == pv1.shape == (144,)
                and load2.shape == pv2.shape == (365, 144)
                and forecasts.shape == (365, 4, 24), {
                    "附件1": list(price.shape), "附件2负荷": list(load2.shape),
                    "附件2光伏": list(pv2.shape), "附件3": list(forecasts.shape),
                })
    audit.check("输入均为有限非负数", all(np.isfinite(a).all() and np.all(a >= 0)
                for a in (price, load1, pv1, load2, pv2, forecasts)), "读取后检查")
    _input_failure_tests(audit, paths["data2"], paths["data3"], dates)

    summaries = _load_summaries(audit)
    if "q1_engineering_baseline" in summaries:
        summary = summaries["q1_engineering_baseline"]
        _constraint_checks(audit, "Q1工程基线", summary)
        _validate_q1_book(audit, SCENARIOS["q1_engineering_baseline"], summary)
        audit.close("Q1工程基线费用回归", summary["objective_cost_cny"], EXPECTED_COSTS["q1"], 0.01)
    for label in ("q2_engineering_baseline", "q2_causal_zero", "q2_causal_baseline",
                  "q2_weekday_optimized"):
        if label in summaries:
            _constraint_checks(audit, label, summaries[label])
            _validate_q2_book(audit, SCENARIOS[label], summaries[label], price, label)
    if "q2_engineering_baseline" in summaries:
        audit.close("Q2工程基线费用回归", summaries["q2_engineering_baseline"]["total_cost_cny"],
                    EXPECTED_COSTS["q2"], 0.01)
    for label in ("q3_update_0", "q3_update_0_6", "q3_update_0_6_12",
                  "q3_update_full_no_terminal", "q3_update_full_terminal"):
        if label in summaries:
            _constraint_checks(audit, label, summaries[label])
            _validate_q3_book(audit, SCENARIOS[label], summaries[label], price, label)
    if "q3_update_full_no_terminal" in summaries:
        audit.close("Q3工程基线费用回归", summaries["q3_update_full_no_terminal"]["total_cost_cny"],
                    EXPECTED_COSTS["q3"], 0.01)
    _stress_test(audit)

    rows = _write_comparison(summaries)
    report = {
        "status": "passed" if audit.passed else "failed",
        "checks_passed": sum(item["passed"] for item in audit.checks),
        "checks_total": len(audit.checks),
        "attachment_sha256": {
            key: file_sha256(paths[key]) for key in ("data1", "data2", "data3", "template1", "template2", "template3")
        },
        "checks": audit.checks,
        "model_comparison": rows,
    }
    output = ROOT / "output"
    output.mkdir(exist_ok=True)
    (output / "validation_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    failed = [item for item in audit.checks if not item["passed"]]
    lines = [
        "# 第一至三问统一验证报告", "",
        f"状态：{'通过' if audit.passed else '失败'}（{report['checks_passed']}/{report['checks_total']}项）", "",
        "## 失败项", "",
    ]
    lines.extend([f"- {item['name']}：{item['evidence']}" for item in failed] or ["- 无"])
    lines.extend(["", "## 验证范围", "",
                  "严格输入读取及四类错误输入失败测试；官方模板结构、日期、逐日合计与费用；",
                  "关键费用回归、SOC和能量平衡约束、极端缺电与光伏过剩压力测试；模型版本汇总。", "",
                  "第四问按当前范围明确排除。", ""])
    (output / "validation_report.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"status": report["status"], "passed": report["checks_passed"],
                      "total": report["checks_total"], "failed": failed}, ensure_ascii=False, indent=2))
    if not audit.passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
