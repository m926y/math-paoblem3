#!/usr/bin/env python3
"""C题问题3：0/6/12/18点滚动时域购电与储能优化。

模型单位：功率 kW，时段电量 kWh，电价 CNY/kWh。
求解器：scipy.optimize.linprog(method="highs")。
"""

from __future__ import annotations

import argparse
import json
import math
from copy import copy
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


DT = 1.0 / 6.0
T = 144
BLOCK = 36
ISSUE_SLOTS = (0, 6, 12, 18)
ISSUE_INDEX = (0, 36, 72, 108)
WEIGHTS = np.array([0.40, 0.25, 0.15, 0.10, 0.05, 0.03, 0.02])
HISTORY_DAYS = 31
RESERVE_FRACTION = 0.20


@dataclass(frozen=True)
class Params:
    eta: float = 0.9
    e_min: float = 1200.0
    e_max: float = 10800.0
    e_initial: float = 6000.0
    power_max: float = 5000.0
    quantile: float = 0.80
    shrink_k: float = 20.0
    cycle_penalty: float = 1e-6
    reserve_fraction: float = RESERVE_FRACTION

    @property
    def q_max(self) -> float:
        return self.power_max * DT


def _to_date(value) -> datetime.date:
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def read_inputs(data1: Path, data2: Path, data3: Path):
    wb1 = load_workbook(data1, read_only=True, data_only=True)
    ws1 = wb1.active
    rows1 = list(ws1.iter_rows(min_row=2, max_row=145, values_only=True))
    price = np.array([float(row[1]) for row in rows1])
    fallback_load = np.array([float(row[2]) for row in rows1])

    wb2 = load_workbook(data2, read_only=True, data_only=True)
    ws_l, ws_v = wb2["小区负载"], wb2["光伏发电实际功率"]
    rows_l = list(ws_l.iter_rows(min_row=2, values_only=True))
    rows_v = list(ws_v.iter_rows(min_row=2, values_only=True))
    dates = [_to_date(row[0]) for row in rows_l]
    load_kw = np.array([[float(x or 0.0) for x in row[1:145]] for row in rows_l])
    pv_kw = np.array([[float(x or 0.0) for x in row[1:145]] for row in rows_v])

    wb3 = load_workbook(data3, read_only=True, data_only=True)
    ws3 = wb3.active
    forecast = np.zeros((len(dates), 4, 24), dtype=float)
    date_to_idx = {d: i for i, d in enumerate(dates)}
    current_date = None
    issue_to_idx = {"0:00": 0, "6:00": 1, "12:00": 2, "18:00": 3}
    for row in ws3.iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            current_date = _to_date(row[0])
        issue = str(row[1]).strip()
        if current_date in date_to_idx and issue in issue_to_idx:
            forecast[date_to_idx[current_date], issue_to_idx[issue], :] = [
                float(x or 0.0) for x in row[2:26]
            ]
    return dates, price, fallback_load, load_kw, pv_kw, forecast


def weighted_past_profile(values: np.ndarray, day: int, fallback: np.ndarray) -> np.ndarray:
    """只用决策日前已完整观测的日期；年初采用递增历史，首日用附件1曲线。"""
    n = min(7, day)
    if n == 0:
        return fallback.copy()
    w = WEIGHTS[:n].copy()
    w /= w.sum()
    return np.sum(values[day - np.arange(1, n + 1)] * w[:, None], axis=0)


def build_load_errors(load_kw: np.ndarray, fallback: np.ndarray) -> np.ndarray:
    errors = np.zeros_like(load_kw)
    for d in range(len(load_kw)):
        pred = weighted_past_profile(load_kw, d, fallback)
        errors[d] = np.maximum(load_kw[d] - pred, 0.0)
    return errors


def shrunken_quantile(local: np.ndarray, global_samples: np.ndarray,
                       q: float, k: float) -> float:
    local = local[np.isfinite(local)]
    global_samples = global_samples[np.isfinite(global_samples)]
    if len(global_samples) == 0:
        return 0.0
    q_global = float(np.quantile(global_samples, q))
    if len(local) == 0:
        return q_global
    q_local = float(np.quantile(local, q))
    alpha = len(local) / (len(local) + k)
    return alpha * q_local + (1.0 - alpha) * q_global


def load_safe_profile(day: int, load_kw: np.ndarray, fallback: np.ndarray,
                      load_errors: np.ndarray, params: Params) -> np.ndarray:
    pred = weighted_past_profile(load_kw, day, fallback)
    history = load_errors[:day]
    if day == 0:
        return np.maximum(pred, 0.0)
    global_hist = history.ravel()
    margin = np.array([
        shrunken_quantile(history[:, t], global_hist, params.quantile, params.shrink_k)
        for t in range(T)
    ])
    return np.maximum(pred + margin, 0.0)


def build_pv_hourly_errors(pv_kw: np.ndarray, forecast: np.ndarray) -> np.ndarray:
    """误差维度为 [发布日期, 发布时间, 提前小时]，仅记录附件覆盖范围内的目标。"""
    days = len(pv_kw)
    flat = pv_kw.ravel()
    errors = np.full((days, 4, 24), np.nan)
    for d in range(days):
        for j, issue_idx in enumerate(ISSUE_INDEX):
            boundary = d * T + issue_idx
            for h in range(24):
                target_end = boundary + 6 * (h + 1) - 1
                if target_end < len(flat):
                    errors[d, j, h] = max(forecast[d, j, h] - flat[target_end], 0.0)
    return errors


def pv_safe_horizon(day: int, issue_j: int, pv_kw: np.ndarray,
                    forecast: np.ndarray, errors: np.ndarray,
                    params: Params) -> np.ndarray:
    issue_idx = ISSUE_INDEX[issue_j]
    boundary = day * T + issue_idx
    flat_actual = pv_kw.ravel()
    anchor = flat_actual[boundary - 1] if boundary > 0 else 0.0
    x_hour = np.arange(0, 25) * 60
    y_hour = np.r_[anchor, forecast[day, issue_j]]
    lead_min = np.arange(1, T + 1) * 10
    interp = np.interp(lead_min, x_hour, y_hour)

    history = errors[:day, issue_j, :]
    global_hist = history.ravel()
    margins = np.zeros(24)
    for h in range(24):
        margins[h] = shrunken_quantile(
            history[:, h] if day else np.array([]), global_hist,
            params.quantile, params.shrink_k
        )
    group = np.minimum((lead_min - 1) // 60, 23).astype(int)
    return np.maximum(interp - margins[group], 0.0)


def reserve_target_soc(price: np.ndarray, load_safe_kw: np.ndarray,
                       pv_safe_kw: np.ndarray, params: Params) -> float:
    """根据下一日高价时段安全净负荷确定跨日储能保留目标。"""
    net_need = np.maximum(load_safe_kw - pv_safe_kw, 0.0) * DT
    expensive = price >= np.quantile(price, 0.75)
    target = params.e_min + params.reserve_fraction * float(np.sum(net_need[expensive]))
    return float(np.clip(target, params.e_min, params.e_max))


def repeat_profile(profile: np.ndarray, start: int, n: int = T) -> np.ndarray:
    return np.array([profile[(start + i) % T] for i in range(n)], dtype=float)


def solve_lp(load_kwh: np.ndarray, pv_kwh: np.ndarray, prices: np.ndarray,
             e0: float, params: Params, reference: np.ndarray | None = None,
             terminal_energy: float | None = None):
    """求解24小时线性规划；reference有限的位置按0点原计划计调整费用。"""
    n = len(load_kwh)
    # x = [g,c,r,s,E,u,v]
    nv = 7 * n
    G, C, R, S, E, U, V = [i * n for i in range(7)]
    adjusted = np.zeros(n, dtype=bool) if reference is None else np.isfinite(reference)
    n_eq = 2 * n + int(adjusted.sum()) + int(terminal_energy is not None)
    aeq = lil_matrix((n_eq, nv), dtype=float)
    beq = np.zeros(n_eq)
    row = 0
    for t in range(n):
        aeq[row, G + t] = 1.0
        aeq[row, R + t] = 1.0
        aeq[row, C + t] = -1.0
        aeq[row, S + t] = -1.0
        beq[row] = load_kwh[t] - pv_kwh[t]
        row += 1
    for t in range(n):
        aeq[row, E + t] = 1.0
        aeq[row, C + t] = -params.eta
        aeq[row, R + t] = 1.0 / params.eta
        if t:
            aeq[row, E + t - 1] = -1.0
            beq[row] = 0.0
        else:
            beq[row] = e0
        row += 1
    if reference is not None:
        for t in np.flatnonzero(adjusted):
            aeq[row, G + t] = 1.0
            aeq[row, U + t] = -1.0
            aeq[row, V + t] = 1.0
            beq[row] = reference[t]
            row += 1
    if terminal_energy is not None:
        aeq[row, E + n - 1] = 1.0
        beq[row] = float(terminal_energy)
        row += 1

    obj = np.zeros(nv)
    obj[G:G+n] = prices
    obj[C:C+n] = params.cycle_penalty
    obj[R:R+n] = params.cycle_penalty
    obj[U:U+n] = 0.5 * prices
    obj[V:V+n] = 0.5 * prices
    bounds = []
    bounds.extend([(0.0, None)] * n)                  # g
    bounds.extend([(0.0, params.q_max)] * n)          # c
    bounds.extend([(0.0, params.q_max)] * n)          # r
    bounds.extend([(0.0, None)] * n)                  # s
    bounds.extend([(params.e_min, params.e_max)] * n) # E
    bounds.extend([(0.0, None) if adjusted[t] else (0.0, 0.0) for t in range(n)])
    bounds.extend([(0.0, None) if adjusted[t] else (0.0, 0.0) for t in range(n)])

    result = linprog(obj, A_eq=aeq.tocsr(), b_eq=beq, bounds=bounds, method="highs")
    if not result.success:
        raise RuntimeError(f"LP求解失败: {result.message}")
    x = result.x
    return {
        "g": x[G:G+n], "c": x[C:C+n], "r": x[R:R+n],
        "s": x[S:S+n], "E": x[E:E+n], "u": x[U:U+n], "v": x[V:V+n],
        "objective": float(result.fun),
    }


def execute_block(g: np.ndarray, charge_plan: np.ndarray, discharge_plan: np.ndarray,
                  load_real: np.ndarray, pv_real: np.ndarray, e0: float,
                  params: Params):
    """按正常购电计划运行，并用实际偏差进行有限储能再调度。"""
    charge_plan = np.asarray(charge_plan, dtype=float)
    discharge_plan = np.asarray(discharge_plan, dtype=float)
    g = np.asarray(g, dtype=float)
    load_real = np.asarray(load_real, dtype=float)
    pv_real = np.asarray(pv_real, dtype=float)
    charge = np.zeros_like(charge_plan)
    discharge = np.zeros_like(discharge_plan)
    emergency = np.zeros_like(g)
    curtail = np.zeros_like(g)
    soc = np.empty(len(g) + 1, dtype=float)
    soc[0] = e0

    for t in range(len(g)):
        energy = float(soc[t])
        c = float(charge_plan[t])
        r = float(discharge_plan[t])
        # 日前计划基于预测SOC，实际SOC可能不同；先把基线动作截断到当前可行范围。
        c = min(c, max((params.e_max - energy) / params.eta, 0.0))
        r = min(r, max((energy - params.e_min) * params.eta, 0.0))
        residual = float(load_real[t] + c - pv_real[t] - r - g[t])
        if residual > 0.0:
            cancelled = min(c, residual)
            c -= cancelled
            residual -= cancelled
            energy_after_baseline = energy + params.eta * c - r / params.eta
            extra = min(
                params.q_max - r,
                residual,
                max((energy_after_baseline - params.e_min) * params.eta, 0.0),
            )
            r += extra
            residual -= extra
        elif residual < 0.0:
            surplus = -residual
            cancelled = min(r, surplus)
            r -= cancelled
            surplus -= cancelled
            extra = min(
                params.q_max - c,
                surplus,
                max((params.e_max - energy - params.eta * c + r / params.eta) / params.eta, 0.0),
            )
            c += extra
            surplus -= extra
            residual = -surplus

        charge[t] = c
        discharge[t] = r
        emergency[t] = max(residual, 0.0)
        curtail[t] = max(-residual, 0.0)
        soc[t + 1] = energy + params.eta * c - r / params.eta
        soc[t + 1] = np.clip(soc[t + 1], params.e_min, params.e_max)

    balance = g + pv_real + discharge + emergency - load_real - charge - curtail
    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "curtail": curtail,
        "soc": soc,
        "balance_residual_max": float(np.max(np.abs(balance), initial=0.0)),
    }


def copy_cell_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def interval_ranges(values: np.ndarray, tol: float = 1e-7):
    active = values > tol
    result = []
    i = 0
    while i < len(values):
        if not active[i]:
            i += 1
            continue
        j = i
        while j + 1 < len(values) and active[j + 1]:
            j += 1
        result.append((i, j, float(values[i:j+1].sum())))
        i = j + 1
    return result


def clock_label(index: int) -> str:
    minute = index * 10
    suffix = "+1" if minute >= 1440 else ""
    minute %= 1440
    return f"{minute // 60}:{minute % 60:02d}{suffix}"


def write_result(template: Path, output: Path, dates, price, records):
    wb = load_workbook(template)
    plan_ws = wb["计划购电量"]
    adj_ws = wb["调整购电量"]
    charge_ws = wb["充放电量"]
    emergency_ws = wb["紧急购电量"]

    out = [r for r in records if r["date"].month >= 2]
    for ws, key in ((plan_ws, "g0"), (adj_ws, "g")):
        for ridx, rec in enumerate(out, start=2):
            ws.cell(ridx, 1, rec["date"])
            ws.cell(ridx, 1).number_format = "yyyy-m-d"
            values = rec[key]
            for t, value in enumerate(values, start=2):
                ws.cell(ridx, t, float(value))
                ws.cell(ridx, t).number_format = "0.0000"
            ws.cell(ridx, 146, float(values.sum()))
            if key == "g0":
                fee = float(np.dot(price, values))
            else:
                fee = float(np.dot(price, values) + 0.5 * np.dot(price, np.abs(values - rec["g0"])))
            ws.cell(ridx, 147, fee)
            ws.cell(ridx, 146).number_format = ws.cell(ridx, 147).number_format = "0.0000"

    # 重建充放电表：每4小时汇总一次；同一行附带日初/日末SOC。
    sample_rows = [charge_ws.cell(r, c) for r in range(2, min(charge_ws.max_row, 7) + 1) for c in range(1, 7)]
    if charge_ws.max_row >= 2:
        charge_ws.delete_rows(2, charge_ws.max_row - 1)
    row = 2
    for rec in out:
        for b in range(6):
            charge_ws.cell(row, 1, rec["date"] if b == 0 else None)
            charge_ws.cell(row, 2, f"{4*b}:00-{4*(b+1)}:00")
            charge_ws.cell(row, 3, float(rec["c"][24*b:24*(b+1)].sum()))
            charge_ws.cell(row, 4, float(rec["r"][24*b:24*(b+1)].sum()))
            if b == 0:
                charge_ws.cell(row, 5, time(0, 0)); charge_ws.cell(row, 6, rec["e0"])
            elif b == 1:
                charge_ws.cell(row, 5, "24:00"); charge_ws.cell(row, 6, rec["E"][-1])
            for cidx in range(1, 7):
                if sample_rows:
                    copy_cell_style(sample_rows[min(b, 5) * 6 + cidx - 1], charge_ws.cell(row, cidx))
            charge_ws.cell(row, 1).number_format = "yyyy-m-d"
            for cidx in (3, 4, 6): charge_ws.cell(row, cidx).number_format = "0.0000"
            row += 1

    # 重建紧急购电表，连续非零十分钟段合并。
    style_src = [copy(emergency_ws.cell(2, c)._style) for c in range(1, 4)] if emergency_ws.max_row >= 2 else []
    if emergency_ws.max_row >= 2:
        emergency_ws.delete_rows(2, emergency_ws.max_row - 1)
    row = 2
    for rec in out:
        for i, j, amount in interval_ranges(rec["emergency"]):
            emergency_ws.cell(row, 1, rec["date"])
            emergency_ws.cell(row, 2, f"{clock_label(i)}-{clock_label(j+1)}")
            emergency_ws.cell(row, 3, amount)
            for cidx in range(1, 4):
                if style_src: emergency_ws.cell(row, cidx)._style = copy(style_src[cidx-1])
            emergency_ws.cell(row, 1).number_format = "yyyy-m-d"
            emergency_ws.cell(row, 3).number_format = "0.0000"
            row += 1
    if row == 2:
        emergency_ws.cell(2, 1, "无紧急购电")

    for ws in wb.worksheets:
        ws.freeze_panes = "B2" if ws in (plan_ws, adj_ws) else "A2"
        ws.auto_filter.ref = ws.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def simulate(dates, price, fallback_load, load_kw, pv_kw, forecast, params: Params):
    load_errors = build_load_errors(load_kw, fallback_load)
    pv_errors = build_pv_hourly_errors(pv_kw, forecast)
    records = []
    e_now = params.e_initial
    january_initial_soc = e_now
    january_warmup_cost = 0.0
    max_balance_residual = 0.0
    max_soc_residual = 0.0

    # 与问题2保持一致：先用1月实际数据进行状态初始化，且根据下一日
    # 安全预测设置终端保留目标，避免1月末因无终端价值而退化到SOC下限。
    for d in range(min(HISTORY_DAYS, len(dates))):
        next_d = min(d + 1, len(dates) - 1)
        next_load_safe = load_safe_profile(next_d, load_kw, fallback_load, load_errors, params)
        next_pv_safe = pv_safe_horizon(next_d, 0, pv_kw, forecast, pv_errors, params)
        target = reserve_target_soc(price, next_load_safe, next_pv_safe, params)
        warm = solve_lp(
            load_kw[d] * DT, pv_kw[d] * DT, price, e_now, params,
            terminal_energy=target,
        )
        january_warmup_cost += float(np.dot(price, warm["g"]))
        e_now = float(warm["E"][-1])
    february_initial_soc = e_now

    for d in range(HISTORY_DAYS, len(dates)):
        date_value = dates[d]
        load_profile = load_safe_profile(d, load_kw, fallback_load, load_errors, params)
        load_h_kw = repeat_profile(load_profile, 0)
        pv_h_kw = pv_safe_horizon(d, 0, pv_kw, forecast, pv_errors, params)
        base = solve_lp(load_h_kw * DT, pv_h_kw * DT, price, e_now, params)
        g0 = base["g"].copy()

        final = {k: np.zeros(T) for k in ("g", "c", "r", "s", "E", "emergency", "curtail")}
        day_e0 = e_now
        day_recourse_energy = 0.0
        for issue_j, start in enumerate(ISSUE_INDEX):
            if issue_j == 0:
                sol = base
            else:
                load_h_kw = repeat_profile(load_profile, start)
                pv_h_kw = pv_safe_horizon(d, issue_j, pv_kw, forecast, pv_errors, params)
                price_h = repeat_profile(price, start)
                reference = np.full(T, np.nan)
                same_day = T - start
                reference[:same_day] = g0[start:]
                sol = solve_lp(load_h_kw * DT, pv_h_kw * DT, price_h, e_now, params, reference)

            take = min(BLOCK, T - start)
            sl = slice(start, start + take)
            actual_load = load_kw[d, sl] * DT
            actual_pv = pv_kw[d, sl] * DT
            executed = execute_block(
                sol["g"][:take], sol["c"][:take], sol["r"][:take],
                actual_load, actual_pv, e_now, params,
            )
            final["g"][sl] = sol["g"][:take]
            final["c"][sl] = executed["charge"]
            final["r"][sl] = executed["discharge"]
            final["s"][sl] = sol["s"][:take]
            final["E"][sl] = executed["soc"][1:]
            final["emergency"][sl] = executed["emergency"]
            final["curtail"][sl] = executed["curtail"]
            max_balance_residual = max(
                max_balance_residual, float(executed["balance_residual_max"])
            )

            prev = executed["soc"][:-1]
            soc_resid = executed["soc"][1:] - prev - params.eta * executed["charge"] + executed["discharge"] / params.eta
            max_soc_residual = max(max_soc_residual, float(np.max(np.abs(soc_resid))))
            day_recourse_energy += float(
                np.sum(np.abs(executed["charge"] - sol["c"][:take]))
                + np.sum(np.abs(executed["discharge"] - sol["r"][:take]))
            )
            e_now = float(executed["soc"][-1])

        records.append({"date": date_value, "e0": day_e0, "g0": g0,
                        "recourse_energy": day_recourse_energy, **final})
        if (d + 1) % 30 == 0 or d + 1 == len(dates):
            print(f"已完成 {d+1}/{len(dates)} 天，当前SOC={e_now:.2f} kWh", flush=True)

    return records, {
        "january_initial_soc": january_initial_soc,
        "february_initial_soc": february_initial_soc,
        "january_warmup_cost_cny": january_warmup_cost,
        "reserve_fraction": params.reserve_fraction,
        "max_actual_balance_residual": max_balance_residual,
        "max_soc_recursion_residual": max_soc_residual,
    }


def summarize(records, price, checks):
    out = [r for r in records if r["date"].month >= 2]
    plan_energy = sum(float(r["g0"].sum()) for r in out)
    final_energy = sum(float(r["g"].sum()) for r in out)
    emergency_energy = sum(float(r["emergency"].sum()) for r in out)
    plan_cost = sum(float(np.dot(price, r["g0"])) for r in out)
    normal_cost = sum(float(np.dot(price, r["g"])) for r in out)
    adjustment_penalty = sum(float(0.5 * np.dot(price, np.abs(r["g"] - r["g0"]))) for r in out)
    emergency_cost = sum(float(5.0 * np.dot(price, r["emergency"])) for r in out)
    simultaneous = max(float(np.max(np.minimum(r["c"], r["r"]))) for r in out)
    summary = {
        "model": "deterministic rolling-horizon MPC-LP",
        "output_dates": [str(out[0]["date"]), str(out[-1]["date"])],
        "output_days": len(out),
        "plan_energy_kwh": plan_energy,
        "final_normal_energy_kwh": final_energy,
        "emergency_energy_kwh": emergency_energy,
        "plan_cost_cny": plan_cost,
        "final_normal_cost_cny": normal_cost,
        "adjustment_penalty_cny": adjustment_penalty,
        "emergency_cost_cny": emergency_cost,
        "total_cost_cny": normal_cost + adjustment_penalty + emergency_cost,
        "soc_min_kwh": min(float(np.min(r["E"])) for r in out),
        "soc_max_kwh": max(float(np.max(r["E"])) for r in out),
        "max_simultaneous_charge_discharge_kwh": simultaneous,
        "actual_charge_energy_kwh": sum(float(r["c"].sum()) for r in out),
        "actual_discharge_energy_kwh": sum(float(r["r"].sum()) for r in out),
        "total_actual_recourse_energy_kwh": sum(float(r["recourse_energy"]) for r in out),
        **checks,
    }
    return summary


def main():
    # q3_model.py所在文件夹
    base_dir = Path(__file__).resolve().parent

    # 附件所在文件夹，与reorganized目录同级
    attachment_dir = base_dir.parent

    # 附件5，即结果模板文件夹
    template_dir = attachment_dir / "附件5"

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data1",
        type=Path,
        default=attachment_dir / "附件1.xlsx"
    )

    parser.add_argument(
        "--data2",
        type=Path,
        default=attachment_dir / "附件2.xlsx"
    )

    parser.add_argument(
        "--data3",
        type=Path,
        default=attachment_dir / "附件3.xlsx"
    )

    parser.add_argument(
        "--template",
        type=Path,
        default=template_dir / "result3.xlsx"
    )

    parser.add_argument(
        "--output",
        type=Path,
        default=base_dir / "result3_计算结果.xlsx"
    )

    parser.add_argument(
        "--summary",
        type=Path,
        default=base_dir / "q3_summary.json"
    )
    parser.add_argument("--quantile", type=float, default=0.80)
    parser.add_argument("--shrink-k", type=float, default=20.0)
    args = parser.parse_args()

    if not 0.5 <= args.quantile < 1.0:
        parser.error("--quantile 应在 [0.5,1) 内")
    params = Params(quantile=args.quantile, shrink_k=args.shrink_k)
    data = read_inputs(args.data1, args.data2, args.data3)
    dates, price, fallback_load, load_kw, pv_kw, forecast = data
    records, checks = simulate(dates, price, fallback_load, load_kw, pv_kw, forecast, params)
    write_result(args.template, args.output, dates, price, records)
    summary = summarize(records, price, checks)
    args.summary.parent.mkdir(parents=True, exist_ok=True)
    args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
