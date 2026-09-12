"""第三问Excel输入输出：读取附件1/2/3并按附件5模板写result3。"""

from __future__ import annotations

from copy import copy
from datetime import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from project_io import read_attachment1, read_attachment2, read_attachment3


T = 144


def read_inputs(data1: Path, data2: Path, data3: Path):
    price, fallback_load, _ = read_attachment1(data1)
    dates, load_kw, pv_kw = read_attachment2(data2)
    forecasts = read_attachment3(data3, dates)
    return dates, price, fallback_load, load_kw, pv_kw, forecasts


def _copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def _clock_label(index: int) -> str:
    minutes = index * 10
    suffix = "+1" if minutes >= 1440 else ""
    minutes %= 1440
    return f"{minutes // 60}:{minutes % 60:02d}{suffix}"


def _active_ranges(values: np.ndarray, tol: float = 1e-7):
    active = np.asarray(values) > tol
    i = 0
    while i < len(values):
        if not active[i]:
            i += 1
            continue
        j = i
        while j + 1 < len(values) and active[j + 1]:
            j += 1
        yield i, j, float(np.sum(values[i:j+1]))
        i = j + 1


def write_result3(template: Path, output: Path, price: np.ndarray, records: list[dict]):
    """保留附件5工作表结构并写入2月1日至12月31日结果。"""

    wb = load_workbook(template)
    required = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if any(name not in wb.sheetnames for name in required):
        raise ValueError(f"result3模板必须包含工作表：{required}")

    output_records = [record for record in records if record["date"].month >= 2]
    plan_ws = wb["计划购电量"]
    adjusted_ws = wb["调整购电量"]
    charge_ws = wb["充放电量"]
    emergency_ws = wb["紧急购电量"]

    for ws, key in ((plan_ws, "grid_plan_0"), (adjusted_ws, "grid_final")):
        for row_no, record in enumerate(output_records, start=2):
            ws.cell(row_no, 1, record["date"])
            ws.cell(row_no, 1).number_format = "yyyy-m-d"
            values = np.asarray(record[key])
            for t, value in enumerate(values, start=2):
                ws.cell(row_no, t, float(value))
                ws.cell(row_no, t).number_format = "0.0000"
            ws.cell(row_no, 146, float(values.sum()))
            if key == "grid_plan_0":
                fee = float(np.dot(price, values))
            else:
                fee = float(
                    np.dot(price, values)
                    + 0.5 * np.dot(price, np.abs(values - record["grid_plan_0"]))
                )
            ws.cell(row_no, 147, fee)
            ws.cell(row_no, 146).number_format = "0.0000"
            ws.cell(row_no, 147).number_format = "0.0000"

    style_rows = [
        [copy(charge_ws.cell(r, c)._style) for c in range(1, 7)]
        for r in range(2, min(charge_ws.max_row, 7) + 1)
    ]
    if charge_ws.max_row >= 2:
        charge_ws.delete_rows(2, charge_ws.max_row - 1)
    row_no = 2
    for record in output_records:
        for block in range(6):
            charge_ws.cell(row_no, 1, record["date"] if block == 0 else None)
            charge_ws.cell(row_no, 2, f"{4*block}:00-{4*(block+1)}:00")
            charge_ws.cell(row_no, 3, float(np.sum(record["charge_actual"][24*block:24*(block+1)])))
            charge_ws.cell(row_no, 4, float(np.sum(record["discharge_actual"][24*block:24*(block+1)])))
            if block == 0:
                charge_ws.cell(row_no, 5, time(0, 0))
                charge_ws.cell(row_no, 6, float(record["soc_initial"]))
            elif block == 1:
                charge_ws.cell(row_no, 5, "24:00")
                charge_ws.cell(row_no, 6, float(record["soc_actual"][-1]))
            if style_rows:
                source = style_rows[min(block, len(style_rows)-1)]
                for column in range(1, 7):
                    charge_ws.cell(row_no, column)._style = copy(source[column-1])
            charge_ws.cell(row_no, 1).number_format = "yyyy-m-d"
            for column in (3, 4, 6):
                charge_ws.cell(row_no, column).number_format = "0.0000"
            row_no += 1

    emergency_style = [copy(emergency_ws.cell(2, c)._style) for c in range(1, 4)]
    if emergency_ws.max_row >= 2:
        emergency_ws.delete_rows(2, emergency_ws.max_row - 1)
    row_no = 2
    for record in output_records:
        for left, right, amount in _active_ranges(record["emergency"]):
            emergency_ws.cell(row_no, 1, record["date"])
            emergency_ws.cell(row_no, 2, f"{_clock_label(left)}-{_clock_label(right+1)}")
            emergency_ws.cell(row_no, 3, amount)
            for column in range(1, 4):
                emergency_ws.cell(row_no, column)._style = copy(emergency_style[column-1])
            emergency_ws.cell(row_no, 1).number_format = "yyyy-m-d"
            emergency_ws.cell(row_no, 3).number_format = "0.0000"
            row_no += 1
    if row_no == 2:
        emergency_ws.cell(2, 1, "无紧急购电")

    for ws in wb.worksheets:
        ws.freeze_panes = "B2" if ws in (plan_ws, adjusted_ws) else "A2"
        ws.auto_filter.ref = ws.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
